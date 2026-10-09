"""CSV / Excel / clipboard exports for e-Invoicing data entry."""
from __future__ import annotations

import base64
import io
import re
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .invoices import Invoice, fmt_date

ITEM_COLUMNS = ["Item Name (Блюдо)", "Quantity (QTY)", "Unit Price (AMD)", "Total Amount (AMD)"]


def _num(x: float):
    """Plain number for files: int when whole, else 2 decimals."""
    x = round(float(x), 2)
    return int(x) if x.is_integer() else x


def safe_filename(*parts: str) -> str:
    s = "_".join(p for p in parts if p)
    s = re.sub(r"[^\w\-.]+", "_", s, flags=re.UNICODE).strip("_")
    return s[:120] or "invoice"


def invoice_frame(inv: Invoice) -> pd.DataFrame:
    return pd.DataFrame({
        ITEM_COLUMNS[0]: inv.items["dish"],
        ITEM_COLUMNS[1]: inv.items["qty"].map(_num),
        ITEM_COLUMNS[2]: inv.items["unit_price"].map(_num),
        ITEM_COLUMNS[3]: inv.items["amount"].map(_num),
    })


def invoice_tsv(inv: Invoice) -> str:
    """Item rows only, tab-separated: paste straight into Excel / an e-Invoicing grid."""
    lines = [
        "\t".join([str(r.dish), str(_num(r.qty)), str(_num(r.unit_price)), str(_num(r.amount))])
        for r in inv.items.itertuples()
    ]
    return "\n".join(lines)


def tsv_b64(inv: Invoice) -> str:
    return base64.b64encode(invoice_tsv(inv).encode("utf-8")).decode("ascii")


def invoice_csv(inv: Invoice, start: date, end: date) -> bytes:
    f = invoice_frame(inv)
    f.insert(0, "Counterparty (Контрагент)", inv.counterparty)
    f.insert(1, "Brand", inv.brand)
    f.insert(2, "Payment Type", ", ".join(inv.pay_types))
    f.insert(3, "Period", _period_text(inv, start, end))
    # utf-8-sig so Excel shows Armenian / Cyrillic correctly on double-click
    return f.to_csv(index=False).encode("utf-8-sig")


def _period_text(inv: Invoice, start: date, end: date) -> str:
    if len(inv.dates) == 1:
        return fmt_date(inv.dates[0])
    return f"{fmt_date(start)} - {fmt_date(end)}"


# ------------------------------------------------------------------ Excel --
_THIN = Side(style="thin", color="E2E5EA")
_HEAD_FILL = PatternFill("solid", fgColor="F3F4F6")
_TOTAL_FILL = PatternFill("solid", fgColor="FFF4E5")


def _fmt(values) -> str:
    return "#,##0" if all(float(v).is_integer() for v in values) else "#,##0.00"


def _write_table(ws, top: int, header: list[str], rows: list[list], num_cols: dict[int, str], widths: list[int]):
    for c, h in enumerate(header, 1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font = Font(bold=True, color="374151")
        cell.fill = _HEAD_FILL
        cell.border = Border(bottom=_THIN, top=_THIN)
        cell.alignment = Alignment(horizontal="right" if c in num_cols else "left", vertical="center")
    for r, row in enumerate(rows, top + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = Border(bottom=_THIN)
            if c in num_cols:
                cell.number_format = num_cols[c]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return top + len(rows) + 1


def invoice_xlsx(inv: Invoice, start: date, end: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = f"Invoice summary — {inv.counterparty}"
    ws["A1"].font = Font(bold=True, size=14)
    meta = [
        ("Counterparty (Контрагент)", inv.counterparty),
        ("Brand", inv.brand),
        ("Payment type", ", ".join(inv.pay_types)),
        ("Period", _period_text(inv, start, end)),
        ("Transfer days", ", ".join(fmt_date(d) for d in inv.dates)),
    ]
    for i, (k, v) in enumerate(meta, 3):
        ws.cell(row=i, column=1, value=k).font = Font(color="6B7280")
        ws.cell(row=i, column=2, value=v).font = Font(bold=True)
    top = 3 + len(meta) + 1
    it = inv.items
    rows = [[r.dish, _num(r.qty), _num(r.unit_price), _num(r.amount)] for r in it.itertuples()]
    nf = {2: _fmt(it["qty"]), 3: _fmt(it["unit_price"]), 4: _fmt(it["amount"])}
    end_row = _write_table(ws, top, ITEM_COLUMNS, rows, nf, [58, 16, 18, 20])
    ws.cell(row=end_row, column=1, value="Total").font = Font(bold=True)
    for c, v in ((2, _num(inv.qty)), (4, _num(inv.amount))):
        cell = ws.cell(row=end_row, column=c, value=v)
        cell.font = Font(bold=True)
        cell.number_format = nf[c]
    for c in range(1, 5):
        ws.cell(row=end_row, column=c).fill = _TOTAL_FILL
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def all_invoices_xlsx(invoices: list[Invoice], start: date, end: date, per_day: bool,
                      brand: str = "", query: str = "") -> bytes:
    """Exactly the cards on screen (brand + period + search) with every invoice line."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Bank transfer invoices"
    ws["A1"].font = Font(bold=True, size=14)
    q = query.strip()
    context = [
        ("Brand", brand or "—"),
        ("Period", f"{fmt_date(start)} – {fmt_date(end)}"),
        ("Search", f"«{q}»" if q else "— (all counterparties)"),
        ("Grouping", "Per day" if per_day else "Whole period"),
    ]
    tot_qty = sum(i.qty for i in invoices)
    tot_amt = sum(i.amount for i in invoices)
    totals = [
        ("Checks / Invoices", len(invoices), "#,##0"),
        ("Items QTY", _num(tot_qty), "#,##0.###" if not float(tot_qty).is_integer() else "#,##0"),
        ("Invoice sales (AMD)", _num(tot_amt), "#,##0"),
    ]
    r = 3
    for k, v in context:
        ws.cell(row=r, column=1, value=k).font = Font(color="6B7280")
        ws.cell(row=r, column=2, value=v).font = Font(bold=True)
        r += 1
    r += 1
    for k, v, f in totals:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=v)
        cell.font = Font(bold=True, size=12)
        cell.number_format = f
        cell.alignment = Alignment(horizontal="left")
        for c in (1, 2):
            ws.cell(row=r, column=c).fill = _TOTAL_FILL
        r += 1
    top = r + 1
    head = ["Counterparty (Контрагент)", "Brand", "Payment Type", "Transfer days",
            "Items QTY", "Invoice Amount (AMD)", "Share of period"]
    rows = [[i.counterparty, i.brand, ", ".join(i.pay_types), ", ".join(fmt_date(d) for d in i.dates),
             _num(i.qty), _num(i.amount), round(i.share, 4)] for i in invoices]
    end_row = _write_table(ws, top, head, rows, {5: _fmt([i.qty for i in invoices] or [0]), 6: "#,##0", 7: "0.0%"},
                           [34, 22, 18, 30, 12, 20, 14])
    ws.cell(row=end_row, column=1, value="Total").font = Font(bold=True)
    for c, v, f in ((5, _num(tot_qty), "#,##0"), (6, _num(tot_amt), "#,##0")):
        cell = ws.cell(row=end_row, column=c, value=v)
        cell.font = Font(bold=True)
        cell.number_format = f
    for c in range(1, 8):
        ws.cell(row=end_row, column=c).fill = _TOTAL_FILL
    ws.freeze_panes = ws.cell(row=top + 1, column=1)

    ws2 = wb.create_sheet("Invoice lines")
    head2 = (["Counterparty (Контрагент)", "Brand"] + (["Date"] if per_day else ["Transfer days"])
             + ITEM_COLUMNS)
    rows2 = []
    for i in invoices:
        when = fmt_date(i.dates[0]) if per_day else ", ".join(fmt_date(d) for d in i.dates)
        for it in i.items.itertuples():
            rows2.append([i.counterparty, i.brand, when, it.dish, _num(it.qty), _num(it.unit_price), _num(it.amount)])
    cols = {c: [row[c - 1] for row in rows2] or [0] for c in (5, 6, 7)}
    end2 = _write_table(ws2, 1, head2, rows2, {c: _fmt(v) for c, v in cols.items()},
                        [30, 12, 22 if not per_day else 12, 58, 15, 17, 19])
    ws2.cell(row=end2, column=1, value=f"Total · {len(rows2)} lines").font = Font(bold=True)
    for c, v in ((5, _num(tot_qty)), (7, _num(tot_amt))):
        cell = ws2.cell(row=end2, column=c, value=v)
        cell.font = Font(bold=True)
        cell.number_format = _fmt(cols[c])
    for c in range(1, 8):
        ws2.cell(row=end2, column=c).fill = _TOTAL_FILL
    ws2.freeze_panes = "A2"
    if rows2:
        ws2.auto_filter.ref = f"A1:G{len(rows2) + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
